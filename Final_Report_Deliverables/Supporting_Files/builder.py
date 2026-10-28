"""
Document formatting and assembly engine for SignBridge AI Final Project Report.
Provides styled helpers for paragraphs, headings, tables, equations, figures, and code blocks.
"""

import os
import docx
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_ALIGN_VERTICAL
from docx.oxml import parse_xml, OxmlElement
from docx.oxml.ns import nsdecls, qn

COLOR_PRIMARY = RGBColor(0x0f, 0x17, 0x2a)      # Slate 900
COLOR_MUTED = RGBColor(0x47, 0x55, 0x69)        # Slate 600
COLOR_HEADER_BG = "1F374C"                     # Midnight Blue Table Header
COLOR_ZEBRA_BG = "F8FAFC"                      # Subtle Table Striping
COLOR_BORDER = "94A3B8"                        # Clean Slate Border

def set_cell_margins(cell, top=100, bottom=100, left=150, right=150):
    tcPr = cell._tc.get_or_add_tcPr()
    tcMar = OxmlElement('w:tcMar')
    for m, val in [('w:top', top), ('w:bottom', bottom), ('w:left', left), ('w:right', right)]:
        node = OxmlElement(m)
        node.set(qn('w:w'), str(val))
        node.set(qn('w:type'), 'dxa')
        tcMar.append(node)
    tcPr.append(tcMar)

def set_cell_shading(cell, color_hex):
    shading_xml = f'<w:shd {nsdecls("w")} w:fill="{color_hex}"/>'
    cell._tc.get_or_add_tcPr().append(parse_xml(shading_xml))

def set_table_borders(table, color="B0BEC5", sz="4", val="single"):
    tblPr = table._tbl.tblPr
    borders_xml = f'''
    <w:tblBorders {nsdecls("w")}>
      <w:top w:val="{val}" w:sz="{sz}" w:space="0" w:color="{color}"/>
      <w:bottom w:val="{val}" w:sz="{sz}" w:space="0" w:color="{color}"/>
      <w:left w:val="none"/>
      <w:right w:val="none"/>
      <w:insideH w:val="{val}" w:sz="{sz}" w:space="0" w:color="{color}"/>
      <w:insideV w:val="none"/>
    </w:tblBorders>
    '''
    tblPr.append(parse_xml(borders_xml))

def add_chapter_heading(doc, chapter_num, chapter_title):
    # Chapter Number
    p_num = doc.add_paragraph()
    p_num.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_num.paragraph_format.space_before = Pt(18)
    p_num.paragraph_format.space_after = Pt(4)
    run_num = p_num.add_run(f"CHAPTER {chapter_num}")
    run_num.font.name = "Times New Roman"
    run_num.font.size = Pt(14)
    run_num.font.bold = True
    run_num.font.color.rgb = COLOR_PRIMARY

    # Chapter Title
    p_title = doc.add_paragraph()
    p_title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_title.paragraph_format.space_before = Pt(0)
    p_title.paragraph_format.space_after = Pt(18)
    run_title = p_title.add_run(chapter_title.upper())
    run_title.font.name = "Times New Roman"
    run_title.font.size = Pt(14)
    run_title.font.bold = True
    run_title.font.color.rgb = COLOR_PRIMARY

def add_heading1(doc, text):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    p.paragraph_format.space_before = Pt(14)
    p.paragraph_format.space_after = Pt(6)
    p.paragraph_format.keep_with_next = True
    run = p.add_run(text)
    run.font.name = "Times New Roman"
    run.font.size = Pt(13)
    run.font.bold = True
    run.font.color.rgb = COLOR_PRIMARY

def add_heading2(doc, text):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    p.paragraph_format.space_before = Pt(10)
    p.paragraph_format.space_after = Pt(4)
    p.paragraph_format.keep_with_next = True
    run = p.add_run(text)
    run.font.name = "Times New Roman"
    run.font.size = Pt(12)
    run.font.bold = True
    run.font.color.rgb = COLOR_PRIMARY

def add_paragraph(doc, text, bold_prefix=None, space_after=6, line_spacing=1.25):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.space_after = Pt(space_after)
    p.paragraph_format.line_spacing = line_spacing

    if bold_prefix:
        r_pre = p.add_run(bold_prefix)
        r_pre.font.name = "Times New Roman"
        r_pre.font.size = Pt(12)
        r_pre.font.bold = True
        r_pre.font.color.rgb = COLOR_PRIMARY

    r = p.add_run(text)
    r.font.name = "Times New Roman"
    r.font.size = Pt(12)
    r.font.color.rgb = COLOR_PRIMARY
    return p

def add_bullet(doc, text, bold_prefix=None, space_after=4):
    p = doc.add_paragraph(style='List Bullet')
    p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.space_after = Pt(space_after)
    p.paragraph_format.line_spacing = 1.15

    if bold_prefix:
        r_pre = p.add_run(bold_prefix)
        r_pre.font.name = "Times New Roman"
        r_pre.font.size = Pt(12)
        r_pre.font.bold = True
        r_pre.font.color.rgb = COLOR_PRIMARY

    r = p.add_run(text)
    r.font.name = "Times New Roman"
    r.font.size = Pt(12)
    r.font.color.rgb = COLOR_PRIMARY

def add_equation(doc, eq_text, eq_num, var_defs=None):
    """
    Renders an equation in a 1-row borderless table:
    Left: centered equation expression
    Right: right-aligned equation number
    Followed by variable definitions.
    """
    table = doc.add_table(rows=1, cols=2)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False

    # 5.3 inches for eq, 0.95 inches for num
    table.columns[0].width = Inches(5.3)
    table.columns[1].width = Inches(0.95)

    cell_eq = table.cell(0, 0)
    cell_num = table.cell(0, 1)

    # Clean borders
    for c in [cell_eq, cell_num]:
        tcPr = c._tc.get_or_add_tcPr()
        borders_xml = f'''
        <w:tcBorders {nsdecls("w")}>
          <w:top w:val="none"/>
          <w:bottom w:val="none"/>
          <w:left w:val="none"/>
          <w:right w:val="none"/>
        </w:tcBorders>
        '''
        tcPr.append(parse_xml(borders_xml))

    p_eq = cell_eq.paragraphs[0]
    p_eq.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_eq.paragraph_format.space_before = Pt(4)
    p_eq.paragraph_format.space_after = Pt(4)
    r_eq = p_eq.add_run(eq_text)
    r_eq.font.name = "Cambria Math"
    r_eq.font.size = Pt(12)
    r_eq.font.italic = True
    r_eq.font.bold = True
    r_eq.font.color.rgb = COLOR_PRIMARY

    p_num = cell_num.paragraphs[0]
    p_num.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    p_num.paragraph_format.space_before = Pt(4)
    p_num.paragraph_format.space_after = Pt(4)
    r_num = p_num.add_run(f"({eq_num})")
    r_num.font.name = "Times New Roman"
    r_num.font.size = Pt(12)
    r_num.font.bold = True
    r_num.font.color.rgb = COLOR_PRIMARY

    # Add variable definitions if provided
    if var_defs:
        p_def = doc.add_paragraph()
        p_def.paragraph_format.space_before = Pt(2)
        p_def.paragraph_format.space_after = Pt(6)
        p_def.paragraph_format.line_spacing = 1.15
        p_def.paragraph_format.left_indent = Inches(0.25)
        r_w = p_def.add_run("where: ")
        r_w.font.name = "Times New Roman"
        r_w.font.size = Pt(11)
        r_w.font.italic = True
        r_w.font.color.rgb = COLOR_MUTED

        for idx, (sym, desc) in enumerate(var_defs):
            if idx > 0:
                p_def = doc.add_paragraph()
                p_def.paragraph_format.space_before = Pt(0)
                p_def.paragraph_format.space_after = Pt(6 if idx == len(var_defs)-1 else 2)
                p_def.paragraph_format.line_spacing = 1.15
                p_def.paragraph_format.left_indent = Inches(0.25)

            rs = p_def.add_run(f"{sym}")
            rs.font.name = "Cambria Math"
            rs.font.size = Pt(11)
            rs.font.italic = True
            rs.font.bold = True
            rs.font.color.rgb = COLOR_PRIMARY

            rd = p_def.add_run(f" = {desc}")
            rd.font.name = "Times New Roman"
            rd.font.size = Pt(11)
            rd.font.color.rgb = COLOR_PRIMARY

def add_figure(doc, img_path, caption_text, width_inches=5.8):
    if not os.path.exists(img_path):
        print(f"Warning: Figure not found at {img_path}")
        return

    p_img = doc.add_paragraph()
    p_img.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_img.paragraph_format.space_before = Pt(8)
    p_img.paragraph_format.space_after = Pt(4)
    run_img = p_img.add_run()
    run_img.add_picture(img_path, width=Inches(width_inches))

    p_cap = doc.add_paragraph()
    p_cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_cap.paragraph_format.space_before = Pt(0)
    p_cap.paragraph_format.space_after = Pt(12)
    p_cap.paragraph_format.keep_with_next = False
    run_cap = p_cap.add_run(caption_text)
    run_cap.font.name = "Times New Roman"
    run_cap.font.size = Pt(11)
    run_cap.font.bold = True
    run_cap.font.color.rgb = COLOR_PRIMARY

def add_table(doc, headers, rows, col_widths=None, caption_text=None):
    if caption_text:
        p_cap = doc.add_paragraph()
        p_cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p_cap.paragraph_format.space_before = Pt(10)
        p_cap.paragraph_format.space_after = Pt(4)
        p_cap.paragraph_format.keep_with_next = True
        run_cap = p_cap.add_run(caption_text)
        run_cap.font.name = "Times New Roman"
        run_cap.font.size = Pt(11)
        run_cap.font.bold = True
        run_cap.font.color.rgb = COLOR_PRIMARY

    num_rows = len(rows) + 1
    num_cols = len(headers)
    table = doc.add_table(rows=num_rows, cols=num_cols)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    set_table_borders(table, color="B0BEC5", sz="4", val="single")

    # Repeat header row across pages
    header_tr = table.rows[0]._tr.get_or_add_trPr()
    header_tr.append(OxmlElement('w:tblHeader'))

    # Format Header Row
    for col_idx, h in enumerate(headers):
        cell = table.cell(0, col_idx)
        cell.text = h
        set_cell_margins(cell, top=120, bottom=120, left=140, right=140)
        set_cell_shading(cell, COLOR_HEADER_BG)
        p = cell.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.space_before = Pt(0)
        p.paragraph_format.space_after = Pt(0)
        for r in p.runs:
            r.font.name = "Times New Roman"
            r.font.size = Pt(10)
            r.font.bold = True
            r.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)

    # Format Data Rows
    for row_idx, row_data in enumerate(rows):
        is_even = (row_idx % 2 == 0)
        row_bg = "FFFFFF" if is_even else COLOR_ZEBRA_BG
        for col_idx, val in enumerate(row_data):
            cell = table.cell(row_idx + 1, col_idx)
            cell.text = str(val)
            set_cell_margins(cell, top=90, bottom=90, left=140, right=140)
            if not is_even:
                set_cell_shading(cell, row_bg)
            p = cell.paragraphs[0]
            # Center if short or number, else left
            val_str = str(val).strip()
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER if (len(val_str) <= 12 or val_str.replace('%','').replace('.','').isdigit()) else WD_ALIGN_PARAGRAPH.LEFT
            p.paragraph_format.space_before = Pt(0)
            p.paragraph_format.space_after = Pt(0)
            p.paragraph_format.line_spacing = 1.15
            for r in p.runs:
                r.font.name = "Times New Roman"
                r.font.size = Pt(10)
                r.font.color.rgb = COLOR_PRIMARY

    # Apply widths if provided
    if col_widths and len(col_widths) == num_cols:
        for row in table.rows:
            for idx, w in enumerate(col_widths):
                row.cells[idx].width = Inches(w)

    p_post = doc.add_paragraph()
    p_post.paragraph_format.space_before = Pt(0)
    p_post.paragraph_format.space_after = Pt(8)

def add_code_block(doc, code_str, title=None):
    if title:
        p_title = doc.add_paragraph()
        p_title.paragraph_format.space_before = Pt(10)
        p_title.paragraph_format.space_after = Pt(3)
        p_title.paragraph_format.keep_with_next = True
        r_t = p_title.add_run(title)
        r_t.font.name = "Times New Roman"
        r_t.font.size = Pt(11)
        r_t.font.bold = True
        r_t.font.color.rgb = COLOR_PRIMARY

    table = doc.add_table(rows=1, cols=1)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    table.columns[0].width = Inches(6.25)

    cell = table.cell(0, 0)
    set_cell_margins(cell, top=140, bottom=140, left=180, right=180)
    set_cell_shading(cell, "F8FAFC")

    # Set subtle border
    tcPr = cell._tc.get_or_add_tcPr()
    borders_xml = f'''
    <w:tcBorders {nsdecls("w")}>
      <w:top w:val="single" w:sz="4" w:space="0" w:color="CBD5E1"/>
      <w:bottom w:val="single" w:sz="4" w:space="0" w:color="CBD5E1"/>
      <w:left w:val="single" w:sz="12" w:space="0" w:color="1F374C"/>
      <w:right w:val="single" w:sz="4" w:space="0" w:color="CBD5E1"/>
    </w:tcBorders>
    '''
    tcPr.append(parse_xml(borders_xml))

    lines = code_str.strip().split("\n")
    cell.paragraphs[0].text = ""  # clear initial
    for idx, l in enumerate(lines):
        p = cell.paragraphs[0] if idx == 0 else cell.add_paragraph()
        p.paragraph_format.space_before = Pt(0)
        p.paragraph_format.space_after = Pt(0)
        p.paragraph_format.line_spacing = 1.05
        r = p.add_run(l)
        r.font.name = "Consolas"
        r.font.size = Pt(9.5)
        r.font.color.rgb = RGBColor(0x1E, 0x29, 0x3B)

    p_post = doc.add_paragraph()
    p_post.paragraph_format.space_before = Pt(0)
    p_post.paragraph_format.space_after = Pt(8)
